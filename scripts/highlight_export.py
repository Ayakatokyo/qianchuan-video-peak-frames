"""Build and export Chapter 04 high-light frame assets as an XLSX workbook."""

from collections import OrderedDict
import os
import posixpath
from pathlib import Path
import tempfile
import xml.etree.ElementTree as ET
import zipfile


CATEGORY_LABELS = {"portrait": "人像画面", "product": "产品画面", "other": "其他画面"}
GALLERY_CATEGORY_ORDER = ("全量", "人像画面", "产品画面", "其他画面")
ASSET_HEADERS = [
    "视觉类型", "素材排名", "素材 ID", "素材标题", "画面类型", "高光帧图片",
    "秒数", "点击峰值排名", "点击指数",
]
GALLERY_COLUMNS = 5
IMAGE_MAX_EDGE_PX = 112
THUMBNAIL_JPEG_QUALITY = 72
IMAGE_MARGIN_PX = 8
IMAGE_COLUMN_WIDTH = 18
PX_TO_POINTS = 0.75


class HighlightExportError(RuntimeError):
    """Raised when the static high-light frame workbook is incomplete."""


def _text(value):
    return str(value).strip() if value is not None else ""


def _category_by_ref(presentation):
    result = {}
    chapter = (presentation or {}).get("chapters", {}).get("chapter04", {})
    for category, group in (chapter.get("groups", {}) or {}).items():
        if not isinstance(group, dict):
            continue
        label = CATEGORY_LABELS.get(category, category)
        for ref in group.get("frameRefs", []):
            if _text(ref):
                result[_text(ref)] = label
    return result


def _frame_ref(material_id, frame):
    return _text(frame.get("frameRef")) or f"highlight:{material_id}:{_text(frame.get('rankType'))}"


def _presentation_material_catalog(data, presentation):
    catalog = {
        _text(material.get("materialId")): material
        for material in (data or {}).get("materials", [])
        if isinstance(material, dict) and _text(material.get("materialId"))
    }
    for material in (presentation or {}).get("materials", []):
        if isinstance(material, dict) and _text(material.get("materialId")):
            catalog[_text(material["materialId"])] = material
    return catalog


def _presentation_frame_facts(data, presentation):
    chapters = (presentation or {}).get("chapters", {})
    chapter05 = chapters.get("chapter05") if isinstance(chapters, dict) else None
    if isinstance(chapter05, dict) and isinstance(chapter05.get("materials"), list):
        source_materials = chapter05["materials"]
    else:
        source_materials = (data or {}).get("materials", [])
    facts = {}
    for material in source_materials:
        if not isinstance(material, dict):
            continue
        material_id = _text(material.get("materialId"))
        highlight = material.get("highlight") if isinstance(material.get("highlight"), dict) else {}
        frames = highlight.get("frames") if isinstance(highlight.get("frames"), list) else material.get("highlightFrames", [])
        for frame in frames:
            if not isinstance(frame, dict):
                continue
            reference = _frame_ref(material_id, frame)
            if reference:
                facts[reference] = (material_id, frame)
    return facts


def build_highlight_export_rows(data, presentation=None):
    """Return user-facing rows from the same frame refs rendered by Chapter 04."""
    category_by_ref = _category_by_ref(presentation)
    catalog = _presentation_material_catalog(data, presentation)
    frame_facts = _presentation_frame_facts(data, presentation)
    chapter04 = (presentation or {}).get("chapters", {}).get("chapter04", {})
    groups = chapter04.get("groups") if isinstance(chapter04, dict) else None
    references = []
    if isinstance(groups, dict):
        for category in ("portrait", "product", "other"):
            group = groups.get(category)
            if not isinstance(group, dict):
                continue
            references.extend(_text(ref) for ref in group.get("frameRefs", []) if _text(ref))
    if not references:
        references = list(frame_facts)

    rows = []
    for reference in references:
        material_id, frame = frame_facts.get(reference, ("", None))
        if not isinstance(frame, dict) or not _text(frame.get("framePath")):
            continue
        material = catalog.get(material_id, {})
        rank_type = _text(frame.get("rankType"))
        peak_rank = rank_type.removeprefix("peak_top_") if rank_type.startswith("peak_top_") else "-"
        frame_label = "封面画面" if rank_type == "zero_second" else f"点击 Top{peak_rank} 波峰画面"
        image_path = str(Path(_text(frame["framePath"])).expanduser().resolve())
        rows.append(OrderedDict([
            ("视觉类型", category_by_ref.get(reference, "其他画面")),
            ("素材排名", material.get("rank")),
            ("素材 ID", material_id),
            ("素材标题", _text(material.get("materialName")) or material_id),
            ("画面类型", frame_label),
            ("秒数", frame.get("durationSec")),
            ("点击峰值排名", peak_rank),
            ("点击指数", frame.get("metricValue")),
            ("__image_path", image_path),
        ]))
    return rows


def _load_workbook_dependencies():
    try:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as ExcelImage
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from PIL import Image
    except ImportError as exc:
        raise HighlightExportError(
            "openpyxl>=3.1,<4 and Pillow>=10,<13 are required for XLSX export"
        ) from exc
    return Workbook, ExcelImage, Alignment, Border, Font, PatternFill, Side, Table, TableStyleInfo, Image


def _image_entry(row, image_class, thumbnail_dir, index):
    image_path = Path(row["__image_path"])
    with image_class.open(image_path) as image:
        image.load()
        thumbnail = image.convert("RGB") if image.mode != "RGB" else image.copy()
    resampling = getattr(image_class, "Resampling", image_class)
    thumbnail.thumbnail(
        (IMAGE_MAX_EDGE_PX, IMAGE_MAX_EDGE_PX),
        getattr(resampling, "LANCZOS", 1),
    )
    thumbnail_path = Path(thumbnail_dir) / f"frame-{index:04d}.jpg"
    thumbnail.save(
        thumbnail_path,
        format="JPEG",
        quality=THUMBNAIL_JPEG_QUALITY,
        optimize=True,
    )
    width, height = thumbnail.size
    thumbnail.close()
    return {
        "row": row,
        "path": thumbnail_path,
        "source_path": image_path,
        "width_px": max(1, width),
        "height_px": max(1, height),
    }


def _make_excel_image(entry, image_class):
    image = image_class(str(entry["path"]))
    image.width = entry["width_px"]
    image.height = entry["height_px"]
    return image


def _style_sheet(sheet, Alignment, Border, Font, PatternFill, Side):
    sheet.sheet_view.showGridLines = False
    border = Border(bottom=Side(style="thin", color="DFCFBD"))
    header_fill = PatternFill("solid", fgColor="2C1810")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    sheet.row_dimensions[1].height = 30


def _build_asset_sheet(workbook, rows, entries, deps):
    _, excel_image_class, Alignment, Border, Font, PatternFill, Side, Table, TableStyleInfo, _ = deps
    sheet = workbook.active
    sheet.title = "高光帧资产"
    sheet.freeze_panes = "A2"
    sheet.append(ASSET_HEADERS)
    body_border = Border(bottom=Side(style="thin", color="DFCFBD"))
    for row_index, row in enumerate(rows, start=2):
        sheet.append([row.get(header) if header != "高光帧图片" else None for header in ASSET_HEADERS])
        sheet.row_dimensions[row_index].height = max(
            96, (entries[row_index - 2]["height_px"] + IMAGE_MARGIN_PX) * PX_TO_POINTS
        )
        for cell in sheet[row_index]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = body_border
        sheet.add_image(_make_excel_image(entries[row_index - 2], excel_image_class), f"F{row_index}")
    widths = {"A": 14, "B": 16, "C": 16, "D": 42, "E": 20, "F": 18, "G": 12, "H": 14, "I": 12}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    _style_sheet(sheet, Alignment, Border, Font, PatternFill, Side)
    table_ref = f"A1:I{max(1, len(rows) + 1)}"
    table = Table(displayName="HighlightFramesTable", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)
    return sheet


def _build_gallery_sheet(workbook, entries, deps):
    _, excel_image_class, Alignment, Border, Font, PatternFill, Side, _, _, _ = deps
    from openpyxl.worksheet.hyperlink import Hyperlink

    sheet = workbook.create_sheet("高光帧图片墙")
    sheet.freeze_panes = "A3"
    sheet.sheet_view.showGridLines = False
    for column in range(1, GALLERY_COLUMNS + 1):
        sheet.column_dimensions[chr(64 + column)].width = IMAGE_COLUMN_WIDTH

    by_category = {label: [] for label in GALLERY_CATEGORY_ORDER[1:]}
    for entry in entries:
        by_category.setdefault(entry["row"].get("视觉类型", "其他画面"), []).append(entry)
    section_rows = {}
    current_row = 3
    section_fill = PatternFill("solid", fgColor="2C1810")
    section_font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    section_border = Border(bottom=Side(style="thin", color="DFCFBD"))
    for label in GALLERY_CATEGORY_ORDER:
        section_rows[label] = current_row
        sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=GALLERY_COLUMNS)
        cell = sheet.cell(current_row, 1, label)
        cell.fill = section_fill
        cell.font = section_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = section_border
        sheet.row_dimensions[current_row].height = 24
        current_row += 1
        section_entries = entries if label == "全量" else by_category.get(label, [])
        for index, entry in enumerate(section_entries):
            column = (index % GALLERY_COLUMNS) + 1
            row = current_row + index // GALLERY_COLUMNS
            sheet.add_image(_make_excel_image(entry, excel_image_class), f"{chr(64 + column)}{row}")
            sheet.row_dimensions[row].height = max(
                sheet.row_dimensions[row].height or 0,
                (entry["height_px"] + IMAGE_MARGIN_PX) * PX_TO_POINTS,
            )
        current_row += max(1, (len(section_entries) + GALLERY_COLUMNS - 1) // GALLERY_COLUMNS) + 1

    nav_fill = PatternFill("solid", fgColor="F7EAD8")
    nav_font = Font(name="Arial", size=11, bold=True, underline="single", color="7A4524")
    nav_border = Border(bottom=Side(style="thin", color="D4B99B"))
    for column, label in enumerate(GALLERY_CATEGORY_ORDER, start=1):
        cell = sheet.cell(1, column, label)
        cell.hyperlink = Hyperlink(
            ref=cell.coordinate,
            location=f"'高光帧图片墙'!A{section_rows[label]}",
        )
        cell.fill = nav_fill
        cell.font = nav_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = nav_border
    sheet.row_dimensions[1].height = 24
    return sheet


def _sheet_drawing_anchors(archive, sheet_number):
    rels_name = f"xl/worksheets/_rels/sheet{sheet_number}.xml.rels"
    if rels_name not in archive.namelist():
        return []
    rels_root = ET.fromstring(archive.read(rels_name))
    drawing_target = next(
        (relationship.get("Target", "") for relationship in rels_root.findall(".//{*}Relationship")
         if relationship.get("Type", "").endswith("/drawing")),
        None,
    )
    if not drawing_target:
        return []
    drawing_name = posixpath.normpath(
        drawing_target.lstrip("/")
        if drawing_target.startswith("/")
        else posixpath.join("xl/worksheets", drawing_target)
    )
    drawing_root = ET.fromstring(archive.read(drawing_name))
    return drawing_root.findall(".//{*}oneCellAnchor")


def _sheet_texts(archive, sheet_number):
    shared_strings = []
    if "xl/sharedStrings.xml" in archive.namelist():
        shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
        for item in shared_root.findall(".//{*}si"):
            shared_strings.append("".join(item.itertext()).strip())
    sheet_root = ET.fromstring(archive.read(f"xl/worksheets/sheet{sheet_number}.xml"))
    texts = []
    for cell in sheet_root.findall(".//{*}c"):
        inline = cell.find("{*}is")
        if inline is not None:
            texts.append("".join(inline.itertext()).strip())
            continue
        value = cell.find("{*}v")
        if value is None or value.text is None:
            continue
        if cell.get("t") == "s":
            index = int(value.text)
            if index < len(shared_strings):
                texts.append(shared_strings[index])
        else:
            texts.append(value.text.strip())
    return texts


def _validate_workbook(path, expected_rows, expected_images, rows):
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
            sheet_names = [sheet.get("name") for sheet in workbook_root.findall(".//{*}sheet")]
            if sheet_names != ["高光帧资产", "高光帧图片墙"]:
                raise HighlightExportError(f"workbook sheets {sheet_names} do not match the required sheets")
            media = {name for name in names if name.startswith("xl/media/") and not name.endswith("/")}
            if len(media) < expected_images:
                raise HighlightExportError(f"workbook image count {len(media)} is below {expected_images} frame rows")
            root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
            sheet_rows = root.findall(".//{*}sheetData/{*}row")
            if len(sheet_rows) != expected_rows + 1:
                raise HighlightExportError(f"workbook row count {len(sheet_rows) - 1} does not match {expected_rows} frame rows")
            if any(row.get("customHeight") != "1" or not row.get("ht") for row in sheet_rows[1:]):
                raise HighlightExportError("asset sheet data rows are missing explicit image-safe heights")
            if len(_sheet_drawing_anchors(archive, 1)) != expected_images:
                raise HighlightExportError("asset sheet image anchors do not match frame rows")
            category_counts = {label: sum(row.get("视觉类型") == label for row in rows) for label in CATEGORY_LABELS.values()}
            expected_gallery_anchors = expected_images + sum(category_counts.values())
            if len(_sheet_drawing_anchors(archive, 2)) != expected_gallery_anchors:
                raise HighlightExportError("gallery image anchors do not match the shared frame rows")
            gallery_texts = _sheet_texts(archive, 2)
            if any(label not in gallery_texts for label in GALLERY_CATEGORY_ORDER):
                raise HighlightExportError("gallery sheet is missing a required section")
            if any(label in gallery_texts for label in ("素材 ID", "点击指数", "素材排名", "秒数")):
                raise HighlightExportError("gallery sheet contains asset metrics")
    except HighlightExportError:
        raise
    except (KeyError, OSError, ValueError, zipfile.BadZipFile, ET.ParseError) as exc:
        raise HighlightExportError(f"unable to validate generated workbook: {path}") from exc


def export_highlight_workbook(data, presentation, output_path):
    rows = build_highlight_export_rows(data, presentation)
    if not rows:
        return None
    deps = _load_workbook_dependencies()
    _, _, _, _, _, _, _, _, _, image_class = deps
    missing_images = [row["__image_path"] for row in rows if not Path(row["__image_path"]).is_file()]
    if missing_images:
        raise HighlightExportError(f"highlight frame image is missing: {missing_images[0]}")
    target = Path(output_path).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="qianchuan-highlight-xlsx-", dir=target.parent) as temp_dir:
        thumbnail_dir = Path(temp_dir) / "thumbnails"
        thumbnail_dir.mkdir()
        try:
            entries = [
                _image_entry(row, image_class, thumbnail_dir, index)
                for index, row in enumerate(rows, start=1)
            ]
        except (image_class.UnidentifiedImageError, OSError) as exc:
            raise HighlightExportError("unable to decode highlight frame image for Excel export") from exc
        temp_output = Path(temp_dir) / target.name
        from openpyxl import Workbook
        workbook = Workbook()
        try:
            _build_asset_sheet(workbook, rows, entries, deps)
            _build_gallery_sheet(workbook, entries, deps)
            workbook.save(temp_output)
        finally:
            workbook.close()
        _validate_workbook(temp_output, len(rows), len(rows), rows)
        os.replace(temp_output, target)
    return target
